import importlib
from io import BytesIO

from django.apps import apps
from django.contrib.auth.models import User
from django.db import connection
from django.test import TestCase, TransactionTestCase
from django.utils import timezone
from openpyxl import Workbook

from src.models import (AuditLog, ResultApproval, ReviewDecision, Section, Session, Student, StudentBehaviouralAssessment,
                        StudentClass, StudentResult, Subject, Term, UserProfile)
from src.scoring import recompute_positions, score_errors


class FixtureMixin:
    def setUp(self):
        self.ss = Section.objects.create(section_name='SS')
        self.ie = Section.objects.create(section_name='IE HQ')
        self.ss1 = StudentClass.objects.create(name='SS 1A', section=self.ss)
        self.ss2 = StudentClass.objects.create(name='SS 2A', section=self.ss)
        self.ie1 = StudentClass.objects.create(name='PRIMARY 1', section=self.ie)
        self.old = StudentClass.objects.create(name='SS 1 2022', section=self.ss, is_archived=True)
        self.session = Session.objects.create(session_name='2025/2026')
        self.term = Term.objects.create(term_name='First Term')
        self.maths = Subject.objects.create(subject_name='Mathematics')
        self.english = Subject.objects.create(subject_name='English Language')
        self.amina = Student.objects.create(student_name='Amina Bello', student_class=self.ss1)
        self.bashir = Student.objects.create(student_name='Bashir Musa', student_class=self.ss1)
        self.ie_kid = Student.objects.create(student_name='Hauwa Sani', student_class=self.ie1)

        self.admin = self.user('admin', UserProfile.ADMIN, None, superuser=True)
        self.exoss = self.user('exoss', UserProfile.EXAMS, self.ss)
        self.exoie = self.user('exoieq', UserProfile.EXAMS, self.ie)
        self.teacher = self.user('teacher', UserProfile.STAFF, self.ss)
        self.head = self.user('hosss', UserProfile.HEAD, self.ss)

    def user(self, username, role, section, superuser=False):
        make = User.objects.create_superuser if superuser else User.objects.create_user
        account = make(username)
        UserProfile.objects.create(user=account, role=role, section=section)
        return account

    def result(self, student, ca1, ca2, exams, cls=None, subject=None):
        return StudentResult.objects.create(
            student=student, session=self.session, term=self.term, student_class=cls or student.student_class,
            subject=subject or self.maths, ca1=ca1, ca2=ca2, exams=exams, total=ca1 + ca2 + exams)

    def behaviour(self, student, cls=None, **scores):
        values = {f: 4 for f in ('conduct', 'punctuality', 'dedication', 'participation',
                                 'hospitality', 'neatness', 'creativity', 'physical')}
        values.update(scores)
        return StudentBehaviouralAssessment.objects.create(
            student=student, session=self.session, term=self.term,
            student_class=cls or student.student_class, **values)


class Fixture(FixtureMixin, TestCase):
    pass


class ScoringTests(TestCase):
    def test_limits(self):
        self.assertEqual(score_errors(20, 20, 60), [])
        self.assertEqual(score_errors('5', 0, '0'), [])
        self.assertEqual(len(score_errors(21, -1, 61)), 3)
        self.assertEqual(len(score_errors('', None, 'abc')), 3)


class AccessTests(Fixture):
    def test_review_requires_exams_officer(self):
        self.assertTemplateUsed(self.client.get('/data-review'), 'src/welcome.html')
        self.client.force_login(self.teacher)
        self.assertEqual(self.client.get('/data-review').status_code, 403)
        self.client.force_login(self.exoss)
        self.assertEqual(self.client.get('/data-review').status_code, 200)

    def test_class_dropdowns_are_scoped_and_hide_archived(self):
        self.client.force_login(self.exoss)
        context = self.client.get('/create-result').context
        self.assertEqual(set(context['classes']), {self.ss1, self.ss2})
        self.assertEqual(list(context['archived_classes']), [self.old])

    def test_unknown_section_sees_no_classes(self):
        nobody = self.user('exth', UserProfile.STAFF, None)
        self.client.force_login(nobody)
        self.assertEqual(list(self.client.get('/create-result').context['classes']), [])


class ScoreReviewTests(Fixture):
    def test_correct_score_recomputes_total_and_positions_and_undo_restores(self):
        bad = self.result(self.amina, 18, 180, 51)
        good = self.result(self.bashir, 10, 10, 40)
        recompute_positions(self.ss1.id, self.maths.id, self.session.id, self.term.id)
        self.client.force_login(self.exoss)
        self.client.post('/data-review/scores', {'result_id': bad.id, 'action': 'save',
                                                 'ca1': '18', 'ca2': '18', 'exams': '51'})
        bad.refresh_from_db()
        good.refresh_from_db()
        self.assertEqual((bad.ca2, bad.total, bad.subject_position), (18, 87, 1))
        self.assertEqual(good.subject_position, 2)

        decision = ReviewDecision.objects.get(kind='result_score')
        self.assertEqual(decision.before['fields']['ca2'], 180)
        self.client.post(f'/data-review/undo/{decision.id}')
        bad.refresh_from_db()
        self.assertEqual((bad.ca2, bad.total), (180, 249))

    def test_rejects_scores_still_out_of_range(self):
        bad = self.result(self.amina, 18, 180, 51)
        self.client.force_login(self.exoss)
        self.client.post('/data-review/scores', {'result_id': bad.id, 'action': 'save',
                                                 'ca1': '18', 'ca2': '25', 'exams': '51'})
        bad.refresh_from_db()
        self.assertEqual(bad.ca2, 180)
        self.assertFalse(ReviewDecision.objects.exists())

    def test_other_section_cannot_touch_result(self):
        bad = self.result(self.amina, 18, 180, 51)
        self.client.force_login(self.exoie)
        response = self.client.get('/data-review/scores')
        self.assertEqual(list(response.context['rows']), [])
        self.client.post('/data-review/scores', {'result_id': bad.id, 'action': 'remove'})
        self.assertTrue(StudentResult.objects.filter(pk=bad.id).exists())

    def test_remove_then_undo_restores_same_row(self):
        bad = self.result(self.amina, 18, 180, 51)
        original = StudentResult.objects.filter(pk=bad.id).values().get()
        original['subject_position'] = 1  # positions are recalculated on restore
        self.client.force_login(self.exoss)
        self.client.post('/data-review/scores', {'result_id': bad.id, 'action': 'remove'})
        self.assertFalse(StudentResult.objects.filter(pk=bad.id).exists())
        decision = ReviewDecision.objects.get()
        self.client.post(f'/data-review/undo/{decision.id}')
        self.assertEqual(StudentResult.objects.filter(pk=bad.id).values().get(), original)

    def test_cannot_undo_twice_or_other_sections_decision(self):
        bad = self.result(self.amina, 18, 180, 51)
        self.client.force_login(self.exoss)
        self.client.post('/data-review/scores', {'result_id': bad.id, 'action': 'remove'})
        decision = ReviewDecision.objects.get()
        self.client.force_login(self.exoie)
        self.client.post(f'/data-review/undo/{decision.id}')
        self.assertFalse(StudentResult.objects.filter(pk=bad.id).exists())
        self.client.force_login(self.exoss)
        self.client.post(f'/data-review/undo/{decision.id}')
        self.client.post(f'/data-review/undo/{decision.id}')
        self.assertEqual(StudentResult.objects.filter(pk=bad.id).count(), 1)


class AcceptAsRecordedTests(Fixture):
    def test_accept_valid_total_but_not_typo(self):
        split = self.result(self.amina, 0, 0, 95)      # whole mark entered as exams
        typo = self.result(self.bashir, 20, 20, 544)
        self.client.force_login(self.exoss)
        self.client.post('/data-review/scores', {'result_id': split.id, 'action': 'accept'})
        self.client.post('/data-review/scores', {'result_id': typo.id, 'action': 'accept'})
        rows = self.client.get('/data-review/scores').context['rows']
        self.assertEqual([r.id for r in rows], [typo.id])
        split.refresh_from_db()
        self.assertEqual((split.ca1, split.ca2, split.exams), (0, 0, 95))

        self.client.post(f'/data-review/undo/{ReviewDecision.objects.get().id}')
        self.assertEqual(len(self.client.get('/data-review/scores').context['rows']), 2)

    def test_accept_batch(self):
        kids = [Student.objects.create(student_name=f'Kid {i}', student_class=self.ss1) for i in range(3)]
        for kid in kids:
            self.result(kid, 0, 0, 80)
        self.result(self.amina, 0, 0, 544)
        self.client.force_login(self.exoss)
        batches = self.client.get('/data-review/scores').context['batches']
        self.assertEqual((batches[0]['count'], batches[0]['acceptable']), (4, 3))
        self.client.post('/data-review/scores', {'action': 'accept_batch', 'class_id': self.ss1.id,
                                                 'session_id': self.session.id, 'term_id': self.term.id})
        self.assertEqual(len(self.client.get('/data-review/scores').context['rows']), 1)

    def test_batch_in_other_section_refused(self):
        for i in range(3):
            self.result(Student.objects.create(student_name=f'Kid {i}', student_class=self.ss1), 0, 0, 80)
        self.client.force_login(self.exoie)
        self.client.post('/data-review/scores', {'action': 'accept_batch', 'class_id': self.ss1.id,
                                                 'session_id': self.session.id, 'term_id': self.term.id})
        self.assertFalse(ReviewDecision.objects.exists())


class DuplicateReviewTests(Fixture):
    def test_keep_one_copy_and_undo(self):
        # Amina's maths result was filed under her class and again under SS 2A.
        right = self.result(self.amina, 10, 10, 40)
        self.result(self.amina, 5, 5, 20, subject=self.english)
        wrong = self.result(self.amina, 1, 1, 1, cls=self.ss2)
        self.client.force_login(self.exoss)
        groups = self.client.get('/data-review/duplicates').context['result_groups']
        self.assertEqual(len(groups), 1)
        self.assertEqual(groups[0]['suggested_id'], right.id)

        self.client.post('/data-review/duplicates', {'kind': 'result', 'key': groups[0]['key'],
                                                     'keep_id': right.id})
        self.assertFalse(StudentResult.objects.filter(pk=wrong.id).exists())
        self.assertTrue(StudentResult.objects.filter(pk=right.id).exists())

        decision = ReviewDecision.objects.get(kind='duplicate_result')
        self.client.post(f'/data-review/undo/{decision.id}')
        self.assertTrue(StudentResult.objects.filter(pk=wrong.id).exists())

    def test_apply_suggested_needs_confirmation(self):
        self.result(self.amina, 10, 10, 40)
        self.result(self.amina, 5, 5, 20, subject=self.english)
        self.result(self.amina, 1, 1, 1, cls=self.ss2)
        self.client.force_login(self.exoss)
        self.client.post('/data-review/duplicates', {'kind': 'result', 'apply_suggested': '1'})
        self.assertEqual(StudentResult.objects.count(), 3)
        self.client.post('/data-review/duplicates', {'kind': 'result', 'apply_suggested': '1', 'confirm': '1'})
        self.assertEqual(StudentResult.objects.count(), 2)

    def test_duplicate_behaviour(self):
        keep = self.behaviour(self.amina)
        drop = self.behaviour(self.amina, cls=self.ss2)
        self.result(self.amina, 10, 10, 40)
        self.client.force_login(self.exoss)
        group = self.client.get('/data-review/duplicates').context['behaviour_groups'][0]
        self.assertEqual(group['suggested_id'], keep.id)
        self.client.post('/data-review/duplicates', {'kind': 'behaviour', 'key': group['key'], 'keep_id': keep.id})
        self.assertFalse(StudentBehaviouralAssessment.objects.filter(pk=drop.id).exists())


class BehaviourReviewTests(Fixture):
    def test_blank_text_value_is_fixed_and_restored_exactly(self):
        row = self.behaviour(self.amina)
        with connection.cursor() as cursor:
            cursor.execute("UPDATE src_studentbehaviouralassessment SET conduct = '' WHERE id = %s", [row.id])
        self.client.force_login(self.exoss)
        self.assertEqual(len(self.client.get('/data-review/behaviour').context['rows']), 1)
        data = {f: '4' for f in ('conduct', 'punctuality', 'dedication', 'participation',
                                 'hospitality', 'neatness', 'creativity', 'physical')}
        self.client.post('/data-review/behaviour', {'row_id': row.id, 'action': 'save', **data})
        self.assertEqual(len(self.client.get('/data-review/behaviour').context['rows']), 0)

        self.client.post(f'/data-review/undo/{ReviewDecision.objects.get().id}')
        with connection.cursor() as cursor:
            cursor.execute('SELECT conduct, typeof(conduct) FROM src_studentbehaviouralassessment WHERE id = %s', [row.id])
            self.assertEqual(cursor.fetchone(), ('', 'text'))


class StudentReviewTests(Fixture):
    def test_remove_empty_duplicate_only(self):
        twin = Student.objects.create(student_name='  amina  bello ', student_class=self.ss1)
        self.result(self.amina, 10, 10, 40)
        self.client.force_login(self.exoss)
        self.assertEqual(len(self.client.get('/data-review/students').context['duplicates']), 1)

        self.client.post('/data-review/students', {'action': 'remove', 'student_id': self.amina.id})
        self.assertTrue(Student.objects.filter(pk=self.amina.id).exists())
        self.client.post('/data-review/students', {'action': 'remove', 'student_id': twin.id})
        self.assertFalse(Student.objects.filter(pk=twin.id).exists())

        self.client.post(f'/data-review/undo/{ReviewDecision.objects.get().id}')
        self.assertEqual(Student.objects.get(pk=twin.id).student_name, '  amina  bello ')

    def test_dismiss_pair(self):
        twin = Student.objects.create(student_name='Amina Bello', student_class=self.ss1)
        self.client.force_login(self.exoss)
        self.client.post('/data-review/students', {'action': 'dismiss', 'pair': f'{self.amina.id}:{twin.id}'})
        self.assertEqual(self.client.get('/data-review/students').context['duplicates'], [])

    def test_only_admin_assigns_class(self):
        lost = Student.objects.create(student_name='Lost Child')
        self.client.force_login(self.exoss)
        self.client.post('/data-review/students', {'action': 'assign', 'student_id': lost.id, 'class_id': self.ss1.id})
        lost.refresh_from_db()
        self.assertIsNone(lost.student_class)
        self.client.force_login(self.admin)
        self.client.post('/data-review/students', {'action': 'assign', 'student_id': lost.id, 'class_id': self.ss1.id})
        lost.refresh_from_db()
        self.assertEqual(lost.student_class, self.ss1)


class PositionTests(Fixture):
    def test_recalculate_and_undo(self):
        a = self.result(self.amina, 10, 10, 40)
        b = self.result(self.bashir, 15, 15, 50)
        StudentResult.objects.update(subject_position=7)
        self.client.force_login(self.exoss)
        self.client.post('/data-review/positions')
        a.refresh_from_db()
        b.refresh_from_db()
        self.assertEqual((a.subject_position, b.subject_position), (2, 1))
        self.client.post(f'/data-review/undo/{ReviewDecision.objects.get().id}')
        self.assertEqual(set(StudentResult.objects.values_list('subject_position', flat=True)), {7})


class EntryValidationTests(FixtureMixin, TransactionTestCase):
    # The legacy upload views commit explicitly, which TestCase's wrapping
    # transaction forbids.

    def upload(self, rows, student_class=None):
        book = Workbook()
        sheet = book.active
        sheet.title = 'result'
        sheet.append(['ID NO', 'Names', '1st C.A', '2nd C.A', 'Exams'])
        for row in rows:
            sheet.append(row)
        buffer = BytesIO()
        book.save(buffer)
        buffer.seek(0)
        buffer.name = 'result.xlsx'
        return self.client.post('/bulk-result-upload', {
            'session': self.session.id, 'term': self.term.id, 'class': (student_class or self.ss1).id,
            'subject': self.maths.id, 'excel_file': buffer})

    def test_out_of_range_score_rejects_whole_sheet(self):
        self.client.force_login(self.exoss)
        self.upload([[self.amina.id, 'Amina', 10, 10, 40], [self.bashir.id, 'Bashir', 10, 10, 544]])
        self.assertFalse(StudentResult.objects.exists())

    def test_student_from_another_class_rejected(self):
        self.client.force_login(self.exoss)
        self.upload([[self.amina.id, 'Amina', 10, 10, 40], [self.ie_kid.id, 'Hauwa', 10, 10, 40]])
        self.assertFalse(StudentResult.objects.exists())

    def test_existing_result_in_other_class_rejected(self):
        self.result(self.amina, 10, 10, 40, cls=self.ss2)
        self.client.force_login(self.exoss)
        self.upload([[self.amina.id, 'Amina', 10, 10, 40]])
        self.assertEqual(StudentResult.objects.count(), 1)

    def test_valid_sheet_with_blank_rows_uploads(self):
        self.client.force_login(self.exoss)
        self.upload([[self.amina.id, 'Amina', 10, 10, 40], [None, None, None, None, None],
                     [self.bashir.id, 'Bashir', 20, 20, 60]])
        totals = dict(StudentResult.objects.values_list('student_id', 'total'))
        self.assertEqual(totals, {self.amina.id: 60, self.bashir.id: 100})


class RestoreMigrationTests(TestCase):
    def test_relinks_only_orphaned_listed_classes(self):
        module = importlib.import_module('src.migrations.0018_restore_np_hq_classes_and_exams_officers')
        np_hq = Section.objects.create(section_name='NP HQ')
        other = Section.objects.create(section_name='SS')
        orphan = StudentClass.objects.create(id=101, name='NURSERY 1A')
        linked = StudentClass.objects.create(id=102, name='NURSERY 1B', section=other)
        unlisted = StudentClass.objects.create(id=999, name='Elsewhere')
        module.restore_np_hq(apps, None)
        orphan.refresh_from_db()
        linked.refresh_from_db()
        unlisted.refresh_from_db()
        self.assertEqual((orphan.section, orphan.is_archived), (np_hq, True))
        self.assertEqual((linked.section, linked.is_archived), (other, False))
        self.assertIsNone(unlisted.section)
        self.assertEqual(ReviewDecision.objects.get().after['class_ids'], [101])


class RolePermissionTests(Fixture):
    def test_staff_cannot_delete_promote_or_manage(self):
        self.client.force_login(self.teacher)
        for url in ('/delete_result', '/promote', '/users', '/audit-trail', '/recycle-bin', '/data-review'):
            self.assertEqual(self.client.get(url).status_code, 403, url)
        self.assertEqual(self.client.get('/create-result').status_code, 200)

    def test_exams_officer_cannot_promote(self):
        self.client.force_login(self.exoss)
        self.assertEqual(self.client.get('/promote').status_code, 403)
        self.assertEqual(self.client.get('/delete_result').status_code, 200)

    def test_tampered_class_from_other_section_is_refused(self):
        self.client.force_login(self.exoie)
        response = self.client.post('/create-result', {
            'form2': '1', 'session': self.session.id, 'term': self.term.id, 'class': self.ss1.id,
            'subject': self.maths.id, 'student_id': [self.amina.id], 'ca1': ['10'], 'ca2': ['10'], 'exams': ['40']})
        self.assertEqual(response.status_code, 403)
        self.assertFalse(StudentResult.objects.exists())

    def test_update_cannot_touch_result_outside_selection(self):
        other = self.result(self.amina, 10, 10, 40, subject=self.english)
        self.client.force_login(self.exoss)
        self.client.post('/update-result', {
            'form2': '1', 'session': self.session.id, 'term': self.term.id, 'class': self.ss1.id,
            'subject': self.maths.id, 'result_id': [other.id], 'ca1': ['1'], 'ca2': ['1'], 'exams': ['1']})
        other.refresh_from_db()
        self.assertEqual(other.ca1, 10)


class EntryFormTests(Fixture):
    def test_on_site_entry_is_self_contained_and_keeps_typed_values_on_error(self):
        self.client.force_login(self.exoss)
        data = {'form2': '1', 'session': self.session.id, 'term': self.term.id, 'class': self.ss1.id,
                'subject': self.maths.id, 'student_id': [self.amina.id, self.bashir.id],
                'ca1': ['15', '25'], 'ca2': ['10', '10'], 'exams': ['50', '40']}
        response = self.client.post('/create-result', data)
        self.assertFalse(StudentResult.objects.exists())
        typed = {s.id: s.typed for s in response.context['students']}
        self.assertEqual(typed[self.bashir.id]['ca1'], '25')
        data['ca1'] = ['15', '20']
        self.client.post('/create-result', data)
        self.assertEqual(dict(StudentResult.objects.values_list('student_id', 'subject_position')),
                         {self.amina.id: 1, self.bashir.id: 2})

    def test_update_records_audit_diff(self):
        result = self.result(self.amina, 10, 10, 40)
        self.client.force_login(self.exoss)
        self.client.post('/update-result', {
            'form2': '1', 'session': self.session.id, 'term': self.term.id, 'class': self.ss1.id,
            'subject': self.maths.id, 'result_id': [result.id], 'ca1': ['12'], 'ca2': ['10'], 'exams': ['40']})
        entry = AuditLog.objects.get(action='update', model='StudentResult')
        self.assertEqual(entry.username, 'exoss')
        self.assertEqual(entry.changes['ca1'], [10, 12])
        self.assertEqual(entry.changes['total'], [60, 62])
        self.assertEqual(entry.section_id, self.ss.id)


class DeleteAndRestoreTests(Fixture):
    def post_delete(self, **extra):
        data = {'form2': '1', 'session': self.session.id, 'term': self.term.id, 'class': self.ss1.id,
                'subject': self.maths.id}
        data.update(extra)
        return self.client.post('/delete_result', data)

    def test_delete_needs_confirmation_reason_and_class_name(self):
        a = self.result(self.amina, 10, 10, 40)
        self.result(self.bashir, 10, 10, 40)
        self.client.force_login(self.exoss)
        response = self.post_delete()
        self.assertEqual(response.context['count'], 2)
        self.post_delete(confirm_delete='1', confirm_class='SS 2A', reason='wrong subject')
        self.post_delete(confirm_delete='1', confirm_class='ss 1a', reason='')
        self.assertEqual(StudentResult.objects.count(), 2)

        self.post_delete(confirm_delete='1', confirm_class='ss 1a', reason='Uploaded under wrong subject')
        self.assertEqual(StudentResult.objects.count(), 0)
        deleted = StudentResult.all_objects.get(pk=a.id)
        self.assertEqual((deleted.deleted_by, deleted.delete_reason), (self.exoss, 'Uploaded under wrong subject'))
        self.assertIsNotNone(deleted.deleted_at)
        self.assertEqual(AuditLog.objects.filter(action='delete', model='StudentResult').count(), 2)
        self.assertTrue(AuditLog.objects.filter(action='bulk_delete').exists())

    def test_recycle_bin_restore_and_conflict(self):
        a = self.result(self.amina, 10, 10, 40)
        self.client.force_login(self.exoss)
        self.post_delete(confirm_delete='1', confirm_class='SS 1A', reason='test')
        group = self.client.get('/recycle-bin').context['groups'][0]
        self.assertEqual(group['count'], 1)

        replacement = self.result(self.amina, 1, 1, 1)
        self.client.post('/recycle-bin', {'kind': group['kind'], 'key': group['key']})
        self.assertIsNotNone(StudentResult.all_objects.get(pk=a.id).deleted_at)

        replacement.soft_delete(self.exoss, reason='made by mistake')
        self.client.post('/recycle-bin', {'kind': group['kind'], 'key': group['key']})
        restored = StudentResult.objects.get(pk=a.id)
        self.assertEqual((restored.ca1, restored.delete_reason), (10, ''))
        self.assertTrue(AuditLog.objects.filter(action='restore', object_id=str(a.id)).exists())

    def test_other_section_cannot_see_or_restore(self):
        self.result(self.amina, 10, 10, 40)
        self.client.force_login(self.exoss)
        self.post_delete(confirm_delete='1', confirm_class='SS 1A', reason='test')
        group = self.client.get('/recycle-bin').context['groups'][0]
        self.client.force_login(self.exoie)
        self.assertEqual(self.client.get('/recycle-bin').context['groups'], [])
        self.client.post('/recycle-bin', {'kind': group['kind'], 'key': group['key']})
        self.assertFalse(StudentResult.objects.exists())

    def test_deleted_results_leave_reports(self):
        self.result(self.amina, 10, 10, 40)
        keep = self.result(self.bashir, 15, 15, 50, subject=self.english)
        self.result(self.amina, 1, 2, 3, subject=self.english).soft_delete(self.exoss)
        self.client.force_login(self.exoss)
        response = self.client.post('/bulk-result-view', {'form1': '1', 'session': self.session.id,
                                                          'term': self.term.id, 'class': self.ss1.id})
        self.assertEqual(response.context['results'][self.amina]['total'], 60)
        self.assertEqual(response.context['results'][self.bashir]['total'], keep.total)

    def test_admin_delete_is_soft(self):
        result = self.result(self.amina, 10, 10, 40)
        self.client.force_login(self.admin)
        self.client.post(f'/admin/src/studentresult/{result.id}/delete/', {'post': 'yes'})
        self.assertEqual(StudentResult.all_objects.get(pk=result.id).delete_reason, 'Deleted in admin panel')

    def test_student_with_results_cannot_be_deleted(self):
        self.result(self.amina, 10, 10, 40)
        self.client.force_login(self.exoss)
        self.client.post('/std-alts', {'delete_items': self.amina.id})
        self.assertTrue(Student.objects.filter(pk=self.amina.id).exists())
        self.client.post('/std-alts', {'delete_items': self.bashir.id})
        self.assertFalse(Student.objects.filter(pk=self.bashir.id).exists())
        self.assertTrue(Student.all_objects.filter(pk=self.bashir.id).exists())


class PromotionTests(Fixture):
    def setUp(self):
        super().setUp()
        self.third = Term.objects.create(term_name='Third Term')
        for student, exams in ((self.amina, 50), (self.bashir, 10)):
            StudentResult.objects.create(student=student, session=self.session, term=self.third,
                                         student_class=self.ss1, subject=self.maths, ca1=5, ca2=5, exams=exams,
                                         total=10 + exams)
        self.carl = Student.objects.create(student_name='Carl New', student_class=self.ss1)

    def choices(self, **extra):
        data = {'session': self.session.id, 'current_class': self.ss1.id, 'promotion_class': self.ss2.id}
        data.update(extra)
        return data

    def test_preview_confirm_and_reverse(self):
        self.client.force_login(self.head)
        rows = {r['student'].id: r for r in self.client.post('/promote', self.choices(preview='1')).context['rows']}
        self.assertTrue(rows[self.amina.id]['promote'])
        self.assertFalse(rows[self.bashir.id]['promote'])
        self.assertIsNone(rows[self.carl.id]['average'])

        self.client.post('/promote', self.choices(confirm='1', student_id=[self.amina.id]))
        self.amina.refresh_from_db()
        self.bashir.refresh_from_db()
        self.assertEqual((self.amina.student_class, self.bashir.student_class), (self.ss2, self.ss1))

        entry = AuditLog.objects.get(action='promote')
        self.client.post('/promote', {'undo': entry.id})
        self.amina.refresh_from_db()
        self.assertEqual(self.amina.student_class, self.ss1)
        self.client.post('/promote', {'undo': entry.id})
        self.assertEqual(AuditLog.objects.filter(action='promote_undo').count(), 1)

    def test_head_of_other_section_refused(self):
        other_head = self.user('hosieq', UserProfile.HEAD, self.ie)
        self.client.force_login(other_head)
        response = self.client.post('/promote', self.choices(confirm='1', student_id=[self.amina.id]))
        self.assertEqual(response.status_code, 403)
        self.amina.refresh_from_db()
        self.assertEqual(self.amina.student_class, self.ss1)


class AccountTests(Fixture):
    def test_admin_creates_personal_account_that_must_change_password(self):
        self.client.force_login(self.admin)
        self.client.post('/users', {'username': 'aisha.bello', 'first_name': 'Aisha', 'last_name': 'Bello',
                                    'email': 'a@example.com', 'role': UserProfile.EXAMS, 'section': self.ss.id,
                                    'password': 'Ramadan-Kareem-2026'})
        aisha = User.objects.get(username='aisha.bello')
        self.assertEqual((aisha.profile.role, aisha.profile.section, aisha.profile.must_change_password),
                         (UserProfile.EXAMS, self.ss, True))

        self.client.force_login(aisha)
        self.assertRedirects(self.client.get('/create-result'), '/password-change')
        self.client.post('/password-change', {'old_password': 'Ramadan-Kareem-2026',
                                              'new_password1': 'Second-Term-Exams-9',
                                              'new_password2': 'Second-Term-Exams-9'})
        self.assertEqual(self.client.get('/create-result').status_code, 200)
        self.assertTrue(AuditLog.objects.filter(model='User', object_id=str(aisha.id),
                                                changes__has_key='password').exists())

    def test_admin_cannot_demote_self_and_staff_cannot_manage(self):
        self.client.force_login(self.admin)
        self.client.post(f'/users/{self.admin.id}', {'role': UserProfile.STAFF, 'section': self.ss.id, 'is_active': '1'})
        self.assertEqual(UserProfile.objects.get(user=self.admin).role, UserProfile.ADMIN)
        self.client.force_login(self.exoss)
        self.assertEqual(self.client.get('/users').status_code, 403)

    def test_login_and_failed_login_are_audited(self):
        self.teacher.set_password('Correct-Horse-77')
        self.teacher.save()
        self.client.post('/', {'username': 'teacher', 'password': 'wrong'})
        self.client.post('/', {'username': 'teacher', 'password': 'Correct-Horse-77'})
        actions = list(AuditLog.objects.filter(action__in=['login', 'login_failed']).values_list('action', flat=True))
        self.assertEqual(sorted(actions), ['login', 'login_failed'])


def signature_data_url():
    import base64
    from PIL import Image, ImageDraw
    image = Image.new('RGBA', (300, 100), (0, 0, 0, 0))
    ImageDraw.Draw(image).line((10, 80, 290, 20), fill=(0, 0, 0, 255), width=4)
    out = BytesIO()
    image.save(out, format='PNG')
    return 'data:image/png;base64,' + base64.b64encode(out.getvalue()).decode()


def sheet_upload(name, header, rows):
    book = Workbook()
    sheet = book.active
    sheet.title = name
    sheet.append(header)
    for row in rows:
        sheet.append(row)
    buffer = BytesIO()
    book.save(buffer)
    buffer.seek(0)
    buffer.name = f'{name}.xlsx'
    return buffer


class ReportTests(Fixture):
    def test_alphabetical_behaviour_by_student_and_card_for_everyone(self):
        zainab = Student.objects.create(student_name='zainab  Abdullahi', student_class=self.ss1)
        for student in (zainab, self.bashir, self.amina):
            self.result(student, 10, 10, 40)
        self.result(self.amina, 5, 5, 20, subject=self.english)
        # Behaviour saved in a different order; Bashir has none.
        self.behaviour(zainab, conduct=1)
        self.behaviour(self.amina, conduct=5)
        self.client.force_login(self.exoss)
        response = self.client.post('/bulk-result-view', {'form1': '1', 'session': self.session.id,
                                                          'term': self.term.id, 'class': self.ss1.id})
        results = response.context['results']
        self.assertEqual([s.student_name for s in results], ['Amina Bello', 'Bashir Musa', 'zainab  Abdullahi'])
        self.assertEqual(results[self.amina]['bhv'].conduct, 5)
        self.assertEqual(results[zainab]['bhv'].conduct, 1)
        self.assertIsNone(results[self.bashir]['bhv'])
        self.assertEqual(response.content.decode().count('Examination Report Sheet'), 3)
        self.assertEqual((results[self.amina]['average'], results[self.amina]['c_pos']), (45.0, '3rd'))
        self.assertEqual(results[self.bashir]['c_pos'], '1st')

    def test_broadsheet_columns_line_up_when_a_subject_is_missing(self):
        self.result(self.amina, 10, 10, 40, subject=self.maths)
        self.result(self.amina, 1, 1, 1, subject=self.english)
        self.result(self.bashir, 2, 2, 2, subject=self.english)
        self.client.force_login(self.exoss)
        response = self.client.post('/class-result-summary', {'form1': '1', 'session': self.session.id,
                                                              'term': self.term.id, 'class': self.ss1.id})
        self.assertEqual(list(response.context['eachSubject']), [self.maths, self.english])
        cells = response.context['results'][self.bashir]['cells']
        self.assertIsNone(cells[0])
        self.assertEqual(cells[1].total, 6)
        self.assertEqual(response.context['term'], 'First Term')
        self.assertEqual(response.context['pcHeadTeacher'], 'Principal')


class SigningTests(Fixture):
    def setUp(self):
        super().setUp()
        self.result(self.amina, 10, 10, 40)
        self.result(self.bashir, 15, 15, 50)
        self.sel = {'session': self.session.id, 'term': self.term.id, 'class': self.ss1.id}
        for user in (self.teacher, self.head):
            self.client.force_login(user)
            self.client.post('/my-signature', {'action': 'draw', 'signature_data': signature_data_url()})

    def approve(self):
        self.client.force_login(self.teacher)
        self.client.post('/sign-results', dict(self.sel, action='teacher'))
        self.client.force_login(self.head)
        self.client.post('/sign-results', dict(self.sel, action='head'))
        return ResultApproval.objects.get(kind='head')

    def test_signature_saved_and_staff_cannot_approve(self):
        self.assertTrue(UserProfile.objects.get(user=self.teacher).signature.startswith('data:image/png;base64,'))
        self.client.force_login(self.teacher)
        self.client.post('/sign-results', dict(self.sel, action='head'))
        self.assertFalse(ResultApproval.objects.filter(kind='head').exists())

    def test_cannot_sign_without_signature(self):
        self.client.force_login(self.exoss)
        self.client.post('/sign-results', dict(self.sel, action='teacher'))
        self.assertFalse(ResultApproval.objects.exists())

    def test_approval_prints_signatures_and_locks_results(self):
        approval = self.approve()
        self.assertEqual(approval.snapshot['students'][str(self.bashir.id)]['total'], 80)
        self.client.force_login(self.exoss)
        page = self.client.post('/bulk-result-view', dict(self.sel, form1='1')).content.decode()
        self.assertEqual(page.count('alt="signature"'), 4)   # 2 cards x (teacher + head)
        self.assertIn('/verify/', page)

        result = StudentResult.objects.get(student=self.amina)
        self.client.post('/update-result', dict(self.sel, form2='1', subject=self.maths.id, result_id=[result.id],
                                                ca1=['1'], ca2=['1'], exams=['1']))
        self.client.post('/delete_result', dict(self.sel, form2='1', subject=self.maths.id, confirm_delete='1',
                                                confirm_class='SS 1A', reason='x'))
        self.client.post('/data-review/scores', {'result_id': result.id, 'action': 'remove'})
        result.refresh_from_db()
        self.assertEqual((result.ca1, result.deleted_at), (10, None))

        self.client.force_login(self.head)
        self.client.post('/sign-results', dict(self.sel, action='revoke', approval_id=approval.id, reason='typo'))
        self.client.force_login(self.exoss)
        self.client.post('/update-result', dict(self.sel, form2='1', subject=self.maths.id, result_id=[result.id],
                                                ca1=['12'], ca2=['10'], exams=['40']))
        result.refresh_from_db()
        self.assertEqual(result.ca1, 12)

    def test_verification_page(self):
        from src.signing import student_token
        approval = self.approve()
        url = f'/verify/{approval.code}/{self.amina.id}/{student_token(approval, self.amina.id)}'
        response = self.client.get(url)
        self.assertTrue(response.context['valid'] and response.context['unchanged'])
        self.assertContains(response, 'Genuine')
        self.assertFalse(self.client.get(f'/verify/{approval.code}/{self.bashir.id}/'
                                         f'{student_token(approval, self.amina.id)}').context['valid'])
        # Reopen, change the score: the printed card no longer matches.
        approval.revoked_at = timezone.now()
        approval.save()
        StudentResult.objects.filter(student=self.amina).update(total=99)
        response = self.client.get(url)
        self.assertFalse(response.context['unchanged'])
        self.assertContains(response, 'withdrawn')


class StudentUploadTests(Fixture):
    def test_preview_then_confirm_with_gender_and_category(self):
        self.client.force_login(self.exoss)
        self.assertEqual(self.client.post('/create-student', {'action': 'template'})['Content-Type'],
                         'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        upload = sheet_upload('students', ['No', 'Name', 'Gender', 'Category'], [
            [1, 'Fatima  Sani', 'F', 'Scholarship'.replace('Scholarship', 'Schoolarship')],
            [2, 'amina bello', 'Female', ''],
            [3, 'Musa Ali', 'm', None]])
        response = self.client.post('/create-student', {'action': 'preview', 'class': self.ss1.id, 'excel_file': upload})
        preview = response.context['preview']
        self.assertEqual([r['include'] for r in preview], [True, False, True])
        self.assertEqual(Student.objects.count(), 3)
        self.client.post('/create-student', {'action': 'confirm', 'token': response.context['token'],
                                             'include': ['0', '2']})
        fatima = Student.objects.get(student_name='Fatima Sani')
        self.assertEqual((fatima.gender, fatima.category, fatima.student_class), ('Female', 'Schoolarship', self.ss1))
        self.assertEqual(Student.objects.get(student_name='Musa Ali').gender, 'Male')
        self.assertEqual(Student.objects.count(), 5)

    def test_bad_gender_rejects_sheet(self):
        self.client.force_login(self.exoss)
        upload = sheet_upload('students', ['No', 'Name', 'Gender'], [[1, 'Ali', 'X']])
        response = self.client.post('/create-student', {'action': 'preview', 'class': self.ss1.id, 'excel_file': upload})
        self.assertNotIn('preview', response.context)
        self.assertEqual(Student.objects.count(), 3)


class ExcelUpdateTests(Fixture):
    def test_download_edit_preview_confirm(self):
        a = self.result(self.amina, 10, 10, 40)
        b = self.result(self.bashir, 15, 15, 50)
        sel = {'session': self.session.id, 'term': self.term.id, 'class': self.ss1.id, 'subject': self.maths.id}
        self.client.force_login(self.exoss)
        download = self.client.post('/update-result-excel', dict(sel, action='download'))
        book = openpyxl_load(download.content)
        rows = list(book['update'].iter_rows(min_row=2, values_only=True))
        self.assertEqual([r[2] for r in rows], ['Amina Bello', 'Bashir Musa'])
        upload = sheet_upload('update', ['Result ID', 'Student ID', 'Name', '1st C.A', '2nd C.A', 'Exams'],
                              [[a.id, self.amina.id, 'Amina', 20, 20, 60], [b.id, self.bashir.id, 'Bashir', 15, 15, 50]])
        response = self.client.post('/update-result-excel', dict(sel, action='preview', excel_file=upload))
        self.assertEqual([c['id'] for c in response.context['preview']], [a.id])
        self.client.post('/update-result-excel', {'action': 'confirm', 'token': response.context['token']})
        a.refresh_from_db()
        b.refresh_from_db()
        self.assertEqual((a.total, a.subject_position, b.subject_position), (100, 1, 2))

    def test_concurrent_change_is_refused(self):
        a = self.result(self.amina, 10, 10, 40)
        sel = {'session': self.session.id, 'term': self.term.id, 'class': self.ss1.id, 'subject': self.maths.id}
        self.client.force_login(self.exoss)
        upload = sheet_upload('update', ['Result ID', 'Student ID', 'Name', '1st C.A', '2nd C.A', 'Exams'],
                              [[a.id, self.amina.id, 'Amina', 20, 20, 60]])
        token = self.client.post('/update-result-excel', dict(sel, action='preview', excel_file=upload)).context['token']
        StudentResult.objects.filter(pk=a.id).update(ca1=11)
        self.client.post('/update-result-excel', {'action': 'confirm', 'token': token})
        a.refresh_from_db()
        self.assertEqual((a.ca1, a.exams), (11, 40))

    def test_wrong_class_sheet_rejected(self):
        other = self.result(self.ie_kid, 10, 10, 40)
        sel = {'session': self.session.id, 'term': self.term.id, 'class': self.ss1.id, 'subject': self.maths.id}
        self.result(self.amina, 10, 10, 40)
        self.client.force_login(self.exoss)
        upload = sheet_upload('update', ['Result ID', 'Student ID', 'Name', '1st C.A', '2nd C.A', 'Exams'],
                              [[other.id, self.ie_kid.id, 'Hauwa', 20, 20, 60]])
        response = self.client.post('/update-result-excel', dict(sel, action='preview', excel_file=upload))
        self.assertNotIn('preview', response.context)


def openpyxl_load(content):
    import openpyxl
    return openpyxl.load_workbook(BytesIO(content))


class PasswordResetAndLockoutTests(Fixture):
    def test_reset_unavailable_without_email(self):
        self.assertTemplateUsed(self.client.get('/password-reset'), 'src/password_reset_unavailable.html')

    def test_reset_by_email(self):
        from django.core import mail
        from django.test import override_settings
        self.teacher.email = 'teacher@example.com'
        self.teacher.set_password('Old-Password-123')   # no reset email for unusable passwords
        self.teacher.save()
        UserProfile.objects.filter(user=self.teacher).update(must_change_password=True)
        with override_settings(EMAIL_HOST='smtp.example.com',
                               EMAIL_BACKEND='django.core.mail.backends.locmem.EmailBackend'):
            self.client.post('/password-reset', {'email': 'teacher@example.com'})
            self.client.post('/password-reset', {'email': 'nobody@example.com'})
        self.assertEqual(len(mail.outbox), 1)
        link = next(line for line in mail.outbox[0].body.splitlines() if '/password-reset/' in line)
        path = link.split('testserver', 1)[1]
        response = self.client.get(path, follow=True)
        self.client.post(response.redirect_chain[-1][0], {'new_password1': 'Fresh-Start-2027',
                                                          'new_password2': 'Fresh-Start-2027'})
        self.teacher.refresh_from_db()
        self.assertTrue(self.teacher.check_password('Fresh-Start-2027'))
        self.assertFalse(UserProfile.objects.get(user=self.teacher).must_change_password)

    def test_lockout_after_five_failures(self):
        self.teacher.set_password('Correct-Horse-77')
        self.teacher.save()
        for _ in range(5):
            self.client.post('/', {'username': 'teacher', 'password': 'wrong'})
        response = self.client.post('/', {'username': 'teacher', 'password': 'Correct-Horse-77'})
        self.assertNotIn('_auth_user_id', self.client.session)
        self.assertContains(response, 'Too many failed attempts')

    def test_custom_404(self):
        self.client.force_login(self.exoss)
        response = self.client.get('/no-such-page')
        self.assertEqual(response.status_code, 404)
        self.assertContains(response, 'Page not found', status_code=404)


class ChangedRowsOnlyTests(Fixture):
    def test_unchanged_flagged_row_does_not_block_update(self):
        legacy = self.result(self.amina, 0, 0, 95)       # historic batch, accepted in Data Review
        fix = self.result(self.bashir, 10, 10, 40)
        self.client.force_login(self.exoss)
        self.client.post('/update-result', {
            'form2': '1', 'session': self.session.id, 'term': self.term.id, 'class': self.ss1.id,
            'subject': self.maths.id, 'result_id': [legacy.id, fix.id],
            'ca1': ['0', '12'], 'ca2': ['0', '10'], 'exams': ['95', '40']})
        fix.refresh_from_db()
        legacy.refresh_from_db()
        self.assertEqual((fix.ca1, legacy.exams), (12, 95))

    def test_shared_account_cannot_sign(self):
        self.result(self.amina, 10, 10, 40)
        UserProfile.objects.filter(user=self.head).update(is_legacy_shared=True, signature='data:image/png;base64,x')
        self.client.force_login(self.head)
        self.client.post('/sign-results', {'session': self.session.id, 'term': self.term.id, 'class': self.ss1.id,
                                           'action': 'head'})
        self.assertFalse(ResultApproval.objects.exists())

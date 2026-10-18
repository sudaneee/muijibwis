"""Login, data entry (on-site and Excel), score sheets, students and deletion.

Every view checks the user's role and that the chosen class is in their
section. Multi-step forms carry their own selection and row ids, so nothing
is shared between requests or users. Writes go through the ORM, so each
change lands in the audit trail; deletes are soft and can be restored.
"""
import re
import uuid
from datetime import timedelta
from io import BytesIO

import openpyxl
from django.contrib import messages
from django.contrib.auth import authenticate, login as auth_login, logout as auth_logout
from django.db import transaction
from django.http import HttpResponse
from django.shortcuts import redirect, render
from django.utils import timezone
from openpyxl import Workbook

from src import audit
from src.access import require, scoped_class, scoped_classes
from src.models import (Session, Student, StudentBehaviouralAssessment, StudentResult, Subject, Term)
from src.reports import (class_result_summary, class_summary, result_view,  # noqa: F401 (routed in urls)
                         single_result_view)
from src.scoring import _as_int, behaviour_errors, expected_positions, recompute_positions, score_errors
from src.signing import lock_message, locked_by

# Column order of the behavioural Excel sheet and on-site form (C..J).
EXCEL_BEHAVIOUR_ORDER = ('conduct', 'punctuality', 'dedication', 'participation',
                         'hospitality', 'creativity', 'physical', 'neatness')


# ------------------------------------------------------------------ helpers

def save_virtual_workbook(book):
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def report_errors(request, errors, limit=15):
    """Show validation problems; nothing is saved when there are any."""
    messages.error(request, f'Nothing was saved: {len(errors)} problem(s) found. Fix them and try again.')
    for error in errors[:limit]:
        messages.error(request, error)
    if len(errors) > limit:
        messages.error(request, f'... and {len(errors) - limit} more.')


def _pick(model, value):
    number = _as_int(value)
    return model.objects.filter(pk=number).first() if number is not None else None


def read_selection(request, need_subject=True, include_archived=False):
    """Session/term/class(/subject) from the form; the class must be in scope.

    Returns (selection, None) or (None, error message).
    """
    data = request.POST
    sel = {'session': _pick(Session, data.get('session')), 'term': _pick(Term, data.get('term'))}
    if need_subject:
        sel['subject'] = _pick(Subject, data.get('subject'))
    missing = [k for k, v in sel.items() if v is None]
    if not str(data.get('class', '')).isdigit():
        missing.append('class')
    if missing:
        return None, f'Choose a {", ".join(missing)}.'
    sel['class'] = scoped_class(request.user, data['class'], include_archived)
    return sel, None


def refuse_if_locked(request, sel):
    """Approved (signed) results cannot be changed until a Head of Section reopens them."""
    approval = locked_by(sel['class'].id, sel['session'].id, sel['term'].id)
    if approval:
        messages.error(request, lock_message(approval))
        return True
    return False


def locked_groups():
    from src.models import ResultApproval
    return set(ResultApproval.objects.filter(kind=ResultApproval.HEAD, revoked_at__isnull=True)
               .values_list('student_class_id', 'session_id', 'term_id'))


def selection_label(sel):
    parts = [sel['class'].name, str(sel.get('subject') or ''), str(sel['session']), str(sel['term'])]
    return ' – '.join(p for p in parts if p)


def result_entry_errors(class_id, subject_id, session_id, term_id, rows):
    """Check new result rows before they are saved.

    `rows` is a list of (label, (student_id, ca1, ca2, exams)). Each student
    must belong to the selected class, appear once, have no existing result
    for this subject and term (in any class), and have scores within limits.
    """
    errors = []
    in_class = set(Student.objects.filter(student_class_id=class_id).values_list('id', flat=True))
    student_ids = [_as_int(r[0]) for _, r in rows]
    existing = set(StudentResult.objects.filter(
        subject_id=subject_id, session_id=session_id, term_id=term_id,
        student_id__in=[i for i in student_ids if i is not None]).values_list('student_id', flat=True))
    seen = set()
    for (label, (sid, ca1, ca2, exams)), number in zip(rows, student_ids):
        if number is None:
            errors.append(f'{label}: student ID "{sid}" is not a number')
            continue
        if number not in in_class:
            errors.append(f'{label}: student ID {number} is not in the selected class')
        if number in seen:
            errors.append(f'{label}: student ID {number} appears more than once')
        if number in existing:
            errors.append(f'{label}: student ID {number} already has a result for this subject and term '
                          '(use Result Update to change it)')
        seen.add(number)
        errors.extend(f'{label}: {e}' for e in score_errors(ca1, ca2, exams))
    return errors


def behaviour_entry_errors(class_id, session_id, term_id, rows):
    """Like result_entry_errors, for (label, (student_id, *eight scores)) rows."""
    errors = []
    in_class = set(Student.objects.filter(student_class_id=class_id).values_list('id', flat=True))
    student_ids = [_as_int(r[0]) for _, r in rows]
    existing = set(StudentBehaviouralAssessment.objects.filter(
        session_id=session_id, term_id=term_id,
        student_id__in=[i for i in student_ids if i is not None]).values_list('student_id', flat=True))
    seen = set()
    for (label, values), number in zip(rows, student_ids):
        if number is None:
            errors.append(f'{label}: student ID "{values[0]}" is not a number')
            continue
        if number not in in_class:
            errors.append(f'{label}: student ID {number} is not in the selected class')
        if number in seen:
            errors.append(f'{label}: student ID {number} appears more than once')
        if number in existing:
            errors.append(f'{label}: student ID {number} already has an assessment for this term')
        seen.add(number)
        scores = dict(zip(EXCEL_BEHAVIOUR_ORDER, values[1:]))
        errors.extend(f'{label}: {e}' for e in behaviour_errors(scores))
    return errors


def alphabetical(items, student=lambda x: x):
    """Sort students (or rows with a student) by name, ignoring case and extra spaces."""
    def key(item):
        s = student(item)
        return (re.sub(r'\s+', ' ', s.student_name or '').strip().lower(), s.id)
    return sorted(items, key=key)


def _non_empty(rows):
    return [(label, r) for label, r in rows if any(v not in (None, '') for v in r)]


def read_sheet(request, sheet_name, columns):
    """Rows (label, values) from the uploaded workbook, header row skipped."""
    upload = request.FILES.get('excel_file')
    if upload is None:
        return None, 'Choose an Excel file to upload.'
    from django.conf import settings
    if upload.size > settings.MAX_UPLOAD_BYTES:
        return None, 'That file is too large (limit 5 MB). Use the downloaded template.'
    try:
        book = openpyxl.load_workbook(upload, data_only=True)
    except Exception:
        return None, 'That file could not be read as an Excel (.xlsx) workbook.'
    if sheet_name not in book.sheetnames:
        return None, f'The workbook has no sheet named "{sheet_name}". Use the downloaded template.'
    rows = []
    for number, row in enumerate(book[sheet_name].iter_rows(min_row=2, values_only=True), start=2):
        rows.append((f'Row {number}', tuple(row[c] if c < len(row) else None for c in columns)))
    return _non_empty(rows), None


def save_results(sel, rows, source):
    """Create results from validated rows, recalculate positions, audit the batch."""
    created = []
    with transaction.atomic():
        for _, (sid, ca1, ca2, exams) in rows:
            ca1, ca2, exams = int(ca1), int(ca2), int(exams)
            result = StudentResult(
                student_id=int(sid), session=sel['session'], term=sel['term'], student_class=sel['class'],
                subject=sel['subject'], ca1=ca1, ca2=ca2, exams=exams, total=ca1 + ca2 + exams)
            result.save()
            created.append(result.id)
        recompute_positions(sel['class'].id, sel['subject'].id, sel['session'].id, sel['term'].id)
        audit.log('upload', model='StudentResult', object_repr=f'{len(created)} results: {selection_label(sel)}',
                  section_id=sel['class'].section_id,
                  details={'source': source, 'class_id': sel['class'].id, 'subject_id': sel['subject'].id,
                           'session_id': sel['session'].id, 'term_id': sel['term'].id, 'result_ids': created})
    return created


def save_assessments(sel, rows, source):
    created = []
    with transaction.atomic():
        for _, values in rows:
            scores = {f: int(v) for f, v in zip(EXCEL_BEHAVIOUR_ORDER, values[1:])}
            row = StudentBehaviouralAssessment(student_id=int(values[0]), session=sel['session'],
                                               term=sel['term'], student_class=sel['class'], **scores)
            row.save()
            created.append(row.id)
        audit.log('upload', model='StudentBehaviouralAssessment',
                  object_repr=f'{len(created)} assessments: {selection_label(sel)}',
                  section_id=sel['class'].section_id,
                  details={'source': source, 'class_id': sel['class'].id, 'session_id': sel['session'].id,
                           'term_id': sel['term'].id, 'assessment_ids': created})
    return created


def typed_students(ids, typed_by_id):
    """Students in the order they were submitted, with what was typed attached."""
    students = {str(s.id): s for s in Student.objects.filter(pk__in=[_as_int(i) for i in ids])}
    ordered = []
    for sid in ids:
        if sid in students:
            students[sid].typed = typed_by_id[sid]
            ordered.append(students[sid])
    return ordered


# ------------------------------------------------------------------ login

LOCKOUT_ATTEMPTS = 5
LOCKOUT_MINUTES = 15


def locked_out(username):
    """Too many recent failed logins for this username (counted from the audit trail)."""
    from src.models import AuditLog
    since = timezone.now() - timedelta(minutes=LOCKOUT_MINUTES)
    failures = AuditLog.objects.filter(action='login_failed', timestamp__gte=since,
                                       details__username__iexact=username)
    last_success = AuditLog.objects.filter(action='login', object_repr__iexact=username,
                                           timestamp__gte=since).order_by('-timestamp').first()
    if last_success:
        failures = failures.filter(timestamp__gt=last_success.timestamp)
    return failures.count() >= LOCKOUT_ATTEMPTS


def home(request):
    if request.method == 'POST':
        username = request.POST.get('username', '').strip()
        if username and locked_out(username):
            messages.error(request, f'Too many failed attempts. Try again in {LOCKOUT_MINUTES} minutes, '
                                    'or ask an administrator to reset your password.')
            return render(request, 'src/welcome.html')
        user = authenticate(request, username=username, password=request.POST.get('password', ''))
        if user is not None:
            auth_login(request, user)
            return redirect('home')
        messages.error(request, 'Incorrect Login-in Details')
        return render(request, 'src/welcome.html')
    if request.user.is_authenticated:
        return render(request, 'src/dashboard.html', dashboard_context(request.user))
    return render(request, 'src/welcome.html')


def dashboard_context(user):
    """Headline figures for the user's section in the most recent session with results."""
    from src.access import has_perm, sees_all_sections, user_section
    from src.models import AuditLog, ResultApproval
    classes = scoped_classes(user)
    class_ids = list(classes.values_list('id', flat=True))
    latest = (Session.objects.filter(studentresult__student_class_id__in=class_ids)
              .order_by('-id').first() if class_ids else None)
    results = StudentResult.objects.filter(student_class_id__in=class_ids, session=latest) if latest else None
    latest_term = (Term.objects.filter(studentresult__in=results).order_by('-id').first()
                   if results is not None else None)
    term_classes = (set(results.filter(term=latest_term).values_list('student_class_id', flat=True))
                    if latest_term else set())
    approved = (set(ResultApproval.objects.filter(student_class_id__in=term_classes, session=latest, term=latest_term,
                                                  kind=ResultApproval.HEAD, revoked_at__isnull=True)
                    .values_list('student_class_id', flat=True)) if latest_term else set())
    activity = AuditLog.objects.exclude(action__in=['login', 'logout', 'login_failed'])
    if has_perm(user, 'view_audit'):
        if not sees_all_sections(user):
            activity = activity.filter(section=user_section(user))
    else:
        activity = activity.filter(user=user)
    return {
        'stats': {
            'students': Student.objects.filter(student_class_id__in=class_ids, status=True).count(),
            'classes': len(class_ids),
            'results': results.count() if results is not None else 0,
            'term_classes': len(term_classes),
            'approved': len(approved),
        },
        'latest_session': latest,
        'latest_term': latest_term,
        'activity': activity.select_related('section')[:8],
    }


def logout(request):
    auth_logout(request)
    return redirect('/')


@require('view_reports')
def side_nav(request):
    return render(request, 'src/index.html')


# ------------------------------------------------------------------ results: Excel

def _upload_results(request, require_empty_subject):
    template = 'src/stdupload.html'
    context = {'subjects': Subject.objects.all()}
    if request.method != 'POST':
        return render(request, template, context)
    sel, error = read_selection(request)
    if error:
        messages.error(request, error)
        return render(request, template, context)
    if refuse_if_locked(request, sel):
        return render(request, template, context)
    if require_empty_subject and StudentResult.objects.filter(
            session=sel['session'], term=sel['term'], student_class=sel['class'], subject=sel['subject']).exists():
        messages.error(request, f'Results for {selection_label(sel)} already exist. Use "Add Missed Result" for '
                                'students who were left out, or Result Update to change scores.')
        return render(request, template, context)
    rows, error = read_sheet(request, 'result', (0, 2, 3, 4))
    if error:
        errors = [error]
    elif not rows:
        errors = ['The sheet has no results.']
    else:
        errors = result_entry_errors(sel['class'].id, sel['subject'].id, sel['session'].id, sel['term'].id, rows)
    if errors:
        report_errors(request, errors)
        return render(request, template, context)
    created = save_results(sel, rows, 'excel')
    messages.info(request, f'Result Uploaded Successfully ({len(created)} students, {selection_label(sel)})')
    return render(request, template, context)


@require('enter_results')
def ResultUpload(request):
    return _upload_results(request, require_empty_subject=True)


@require('enter_results')
def not_uploaded_results(request):
    return _upload_results(request, require_empty_subject=False)


# ------------------------------------------------------------------ results: on-site entry

@require('enter_results')
def resultcreate(request):
    context = {'subjects': Subject.objects.all()}
    if request.method == 'POST' and 'form1' in request.POST:
        sel, error = read_selection(request)
        if error:
            messages.error(request, error)
            return render(request, 'src/resultcreate.html', context)
        context.update(sel=sel, students=alphabetical(Student.objects.filter(student_class=sel['class'], status=True)))
        return render(request, 'src/resultcreate.html', context)

    if request.method == 'POST' and 'form2' in request.POST:
        sel, error = read_selection(request)
        if error:
            messages.error(request, error)
            return render(request, 'src/resultcreate.html', context)
        if refuse_if_locked(request, sel):
            return render(request, 'src/resultcreate.html', context)
        ids = request.POST.getlist('student_id')
        ca1s, ca2s, exams = (request.POST.getlist(f) for f in ('ca1', 'ca2', 'exams'))
        names = request.POST.getlist('name') + [''] * len(ids)
        rows = [(name or f'Student ID {sid}', (sid, a, b, e))
                for sid, name, a, b, e in zip(ids, names, ca1s, ca2s, exams)]
        errors = result_entry_errors(sel['class'].id, sel['subject'].id, sel['session'].id, sel['term'].id, rows)
        if errors:
            report_errors(request, errors)
            typed = {sid: {'ca1': a, 'ca2': b, 'exams': e} for sid, a, b, e in zip(ids, ca1s, ca2s, exams)}
            context.update(sel=sel, students=typed_students(ids, typed), show_form=True)
            return render(request, 'src/resultcreate.html', context)
        created = save_results(sel, rows, 'form')
        messages.info(request, f'Result Uploaded Successfully ({len(created)} students, {selection_label(sel)})')
        return render(request, 'src/resultcreate.html', context)

    return render(request, 'src/resultcreate.html', context)


@require('enter_results')
def updateresult(request):
    context = {'subjects': Subject.objects.all()}
    if request.method == 'POST' and 'form1' in request.POST:
        sel, error = read_selection(request)
        if error:
            messages.error(request, error)
            return render(request, 'src/updateresult.html', context)
        results = alphabetical(StudentResult.objects.filter(session=sel['session'], term=sel['term'],
                                                            student_class=sel['class'], subject=sel['subject'])
                               .select_related('student'), student=lambda r: r.student)
        if not results:
            messages.error(request, f'No results found for {selection_label(sel)}.')
        context.update(sel=sel, all_results=results)
        return render(request, 'src/updateresult.html', context)

    if request.method == 'POST' and 'form2' in request.POST:
        sel, error = read_selection(request)
        if error:
            messages.error(request, error)
            return render(request, 'src/updateresult.html', context)
        if refuse_if_locked(request, sel):
            return render(request, 'src/updateresult.html', context)
        ids = [_as_int(i) for i in request.POST.getlist('result_id')]
        ca1s, ca2s, exams = (request.POST.getlist(f) for f in ('ca1', 'ca2', 'exams'))
        results = {r.id: r for r in StudentResult.objects.filter(
            pk__in=[i for i in ids if i is not None], session=sel['session'], term=sel['term'],
            student_class=sel['class'], subject=sel['subject']).select_related('student')}
        errors = []
        for rid, a, b, e in zip(ids, ca1s, ca2s, exams):
            result = results.get(rid)
            if result is None:
                errors.append(f'Result {rid} does not belong to {selection_label(sel)}')
                continue
            if [_as_int(a), _as_int(b), _as_int(e)] == [result.ca1, result.ca2, result.exams]:
                continue   # unchanged rows are left alone, even if Data Review has flagged them
            errors.extend(f'{result.student}: {msg}' for msg in score_errors(a, b, e))
        if errors:
            report_errors(request, errors)
            for rid, a, b, e in zip(ids, ca1s, ca2s, exams):
                if rid in results:
                    results[rid].ca1, results[rid].ca2, results[rid].exams = a, b, e
            context.update(sel=sel, all_results=[results[i] for i in ids if i in results], show_form=True)
            return render(request, 'src/updateresult.html', context)
        changed = 0
        with transaction.atomic():
            for rid, a, b, e in zip(ids, ca1s, ca2s, exams):
                result = results[rid]
                new = (_as_int(a), _as_int(b), _as_int(e))
                if (result.ca1, result.ca2, result.exams) != new:
                    result.ca1, result.ca2, result.exams = new
                    result.total = sum(new)
                    result.save()
                    changed += 1
            recompute_positions(sel['class'].id, sel['subject'].id, sel['session'].id, sel['term'].id)
        messages.info(request, f'Result Updated Successfully ({changed} changed, {selection_label(sel)})')
        return render(request, 'src/updateresult.html', context)

    return render(request, 'src/updateresult.html', context)


@require('enter_results')
def single_result_update(request):
    context = {'subjects': Subject.objects.all()}
    if request.method == 'POST' and 'form1' in request.POST:
        sel, error = read_selection(request)
        if error:
            messages.error(request, error)
            return render(request, 'src/single_result_update.html', context)
        result = StudentResult.objects.filter(
            student=_as_int(request.POST.get('student_id')), session=sel['session'], term=sel['term'],
            student_class=sel['class'], subject=sel['subject']).select_related('student').first()
        if result is None:
            messages.error(request, f'No result for that student ID in {selection_label(sel)}.')
            return render(request, 'src/single_result_update.html', context)
        context.update(sel=sel, get_student=result, student_id=result.student_id)
        return render(request, 'src/single_result_update.html', context)

    if request.method == 'POST' and 'form2' in request.POST:
        sel, error = read_selection(request)
        if error:
            messages.error(request, error)
            return render(request, 'src/single_result_update.html', context)
        if refuse_if_locked(request, sel):
            return render(request, 'src/single_result_update.html', context)
        result = StudentResult.objects.filter(
            pk=_as_int(request.POST.get('result_id')), session=sel['session'], term=sel['term'],
            student_class=sel['class'], subject=sel['subject']).select_related('student').first()
        if result is None:
            messages.error(request, 'That result was not found.')
            return render(request, 'src/single_result_update.html', context)
        ca1, ca2, exams = (request.POST.get(f, '') for f in ('ca1', 'ca2', 'exams'))
        errors = score_errors(ca1, ca2, exams)
        if errors:
            report_errors(request, errors)
            result.ca1, result.ca2, result.exams = ca1, ca2, exams
            context.update(sel=sel, get_student=result, student_id=result.student_id, show_form=True)
            return render(request, 'src/single_result_update.html', context)
        with transaction.atomic():
            result.ca1, result.ca2, result.exams = int(ca1), int(ca2), int(exams)
            result.total = result.ca1 + result.ca2 + result.exams
            result.save()
            recompute_positions(sel['class'].id, sel['subject'].id, sel['session'].id, sel['term'].id)
        messages.info(request, f'Result updated for {result.student}.')
        return render(request, 'src/single_result_update.html', context)

    return render(request, 'src/single_result_update.html', context)


@require('review_data')
def setting_position(request):
    ids = set(scoped_classes(request.user, include_archived=True).values_list('id', flat=True))
    rows = list(StudentResult.objects.filter(student_class_id__in=ids).only(
        'id', 'student_class_id', 'subject_id', 'session_id', 'term_id', 'total', 'subject_position'))
    groups = {}
    for r in rows:
        groups.setdefault((r.student_class_id, r.subject_id, r.session_id, r.term_id), []).append(r)
    changed = []
    locked = locked_groups()
    for key, members in groups.items():
        if (key[0], key[2], key[3]) in locked:
            continue
        expected = expected_positions(members)
        for r in members:
            if r.subject_position != expected[r.id]:
                r.subject_position = expected[r.id]
                changed.append(r)
    StudentResult.objects.bulk_update(changed, ['subject_position'], batch_size=500)
    audit.log('positions', model='StudentResult', object_repr=f'{len(changed)} subject positions recalculated',
              details={'changed': len(changed)})
    messages.info(request, f'Position Set Successfully ({len(changed)} corrected)')
    return redirect('home')


# ------------------------------------------------------------------ behaviour

@require('enter_results')
def settings(request):
    """On-site behavioural assessment entry (the URL name predates this page)."""
    if request.method == 'POST' and 'form1' in request.POST:
        sel, error = read_selection(request, need_subject=False)
        if error:
            messages.error(request, error)
            return render(request, 'src/settings.html', {})
        return render(request, 'src/settings.html', {
            'sel': sel, 'students': alphabetical(Student.objects.filter(student_class=sel['class']))})

    if request.method == 'POST' and 'form2' in request.POST:
        sel, error = read_selection(request, need_subject=False)
        if error:
            messages.error(request, error)
            return render(request, 'src/settings.html', {})
        if refuse_if_locked(request, sel):
            return render(request, 'src/settings.html', {})
        ids = request.POST.getlist('student_id')
        columns = [request.POST.getlist(f) for f in EXCEL_BEHAVIOUR_ORDER]
        rows = [(f'Student ID {sid}', (sid, *vals)) for sid, *vals in zip(ids, *columns)]
        if StudentBehaviouralAssessment.objects.filter(
                session=sel['session'], term=sel['term'], student_class=sel['class']).exists():
            errors = [f'Assessments for {selection_label(sel)} already exist. Use "Upload missed Assessment" '
                      'for students who were left out.']
        else:
            errors = behaviour_entry_errors(sel['class'].id, sel['session'].id, sel['term'].id, rows)
        if errors:
            report_errors(request, errors)
            typed = {sid: dict(zip(EXCEL_BEHAVIOUR_ORDER, vals)) for sid, *vals in zip(ids, *columns)}
            return render(request, 'src/settings.html', {
                'sel': sel, 'students': typed_students(ids, typed), 'show_form': True})
        created = save_assessments(sel, rows, 'form')
        messages.info(request, f'Assessments saved ({len(created)} students, {selection_label(sel)})')
        return render(request, 'src/settings.html', {})

    return render(request, 'src/settings.html', {})


def _upload_behaviour(request, require_empty_class):
    if request.method != 'POST':
        return render(request, 'src/behav.html', {})
    sel, error = read_selection(request, need_subject=False)
    if error:
        messages.error(request, error)
        return render(request, 'src/behav.html', {})
    if refuse_if_locked(request, sel):
        return render(request, 'src/behav.html', {})
    if require_empty_class and StudentBehaviouralAssessment.objects.filter(
            session=sel['session'], term=sel['term'], student_class=sel['class']).exists():
        messages.error(request, f'Assessments for {selection_label(sel)} already exist. Use "Upload missed '
                                'Assessment" for students who were left out.')
        return render(request, 'src/behav.html', {})
    rows, error = read_sheet(request, 'result', (0, 2, 3, 4, 5, 6, 7, 8, 9))
    if error:
        errors = [error]
    elif not rows:
        errors = ['The sheet has no assessments.']
    else:
        errors = behaviour_entry_errors(sel['class'].id, sel['session'].id, sel['term'].id, rows)
    if errors:
        report_errors(request, errors)
        return render(request, 'src/behav.html', {})
    created = save_assessments(sel, rows, 'excel')
    messages.info(request, f'Assessments uploaded ({len(created)} students, {selection_label(sel)})')
    return render(request, 'src/behav.html', {})


@require('enter_results')
def bhvxl(request):
    return _upload_behaviour(request, require_empty_class=True)


@require('enter_results')
def bhvxl_missed(request):
    return _upload_behaviour(request, require_empty_class=False)


# ------------------------------------------------------------------ score sheets

@require('view_reports')
def score_sheet(request):
    if request.method == 'POST' and 'form1' in request.POST:
        student_class = scoped_class(request.user, request.POST.get('class'))
        return render(request, 'src/score_sheet1.html', {
            'classs': str(student_class.id), 'all': alphabetical(Student.objects.filter(student_class=student_class))})
    return render(request, 'src/score_sheet.html', {})


@require('enter_results')
def excel_score_sheet(request):
    if request.method == 'POST' and 'form1' in request.POST:
        student_class = scoped_class(request.user, request.POST.get('class'))
        subject = _pick(Subject, request.POST.get('subject'))
        if subject is None:
            messages.error(request, 'Choose a subject.')
            return render(request, 'src/score_sheet.html', {'subjects': Subject.objects.all()})
        book = Workbook()
        sheet = book.active
        sheet.title = 'result'
        sheet.append(['ID NO', 'Names', '1st C.A', '2nd C.A', 'Exams'])
        for student in alphabetical(Student.objects.filter(student_class=student_class)):
            sheet.append([student.id, student.student_name, 0, 0, 0])
        response = HttpResponse(content=save_virtual_workbook(book), content_type='application/ms-excel')
        response['Content-Disposition'] = f'attachment; filename="{student_class.id}-{subject.id}.xlsx"'
        return response
    return render(request, 'src/score_sheet.html', {'subjects': Subject.objects.all()})


@require('enter_results')
def behav_sheet(request):
    if request.method == 'POST' and 'form1' in request.POST:
        student_class = scoped_class(request.user, request.POST.get('class'))
        book = Workbook()
        sheet = book.active
        sheet.title = 'result'
        sheet.append(['ID NO', 'Names', 'conduct', 'punc.', 'dedic.', 'part.', 'hosp.', 'creat.', 'phy.', 'neat.'])
        for student in alphabetical(Student.objects.filter(student_class=student_class)):
            sheet.append([student.id, student.student_name] + [0] * 8)
        response = HttpResponse(content=save_virtual_workbook(book), content_type='application/ms-excel')
        response['Content-Disposition'] = 'attachment; filename=template.xlsx'
        return response
    return render(request, 'src/behavsheet.html', {})


# ------------------------------------------------------------------ students

@require('manage_students')
def std_alts(request):
    def class_list(student_class):
        return render(request, 'src/std-alts.html',
                      {'student': alphabetical(Student.objects.filter(student_class=student_class))})

    if request.method != 'POST':
        return render(request, 'src/std-alts.html', {})

    if 'form1' in request.POST:
        return class_list(scoped_class(request.user, request.POST.get('class')))

    if 'add' in request.POST:
        student_class = scoped_class(request.user, request.POST.get('class'))
        name = request.POST.get('name', '').strip()
        if not name:
            messages.error(request, 'Enter the student name.')
            return render(request, 'src/std-alts.html', {})
        Student.objects.create(student_name=name, student_class=student_class, status=True, category='Normal')
        messages.success(request, 'Record Created Successfully')
        return render(request, 'src/std-alts.html', {})

    if 'save' in request.POST:
        messages.info(request, 'Promotion has moved to "Promote Students".')
        return redirect('promote')

    if 'delete_items' in request.POST:
        student = Student.objects.filter(pk=_as_int(request.POST.get('delete_items'))).first()
        if student is None:
            messages.error(request, 'That student was not found.')
            return render(request, 'src/std-alts.html', {})
        student_class = scoped_class(request.user, student.student_class_id)
        has_records = (StudentResult.all_objects.filter(student=student).exists()
                       or StudentBehaviouralAssessment.all_objects.filter(student=student).exists())
        if has_records:
            messages.error(request, f'{student} has results or assessments recorded and cannot be deleted. '
                                    'Mark them inactive instead.')
        else:
            student.soft_delete(request.user, reason='Deleted from Other Settings')
            messages.success(request, f'{student} deleted. The record can be restored from the Recycle Bin.')
        return class_list(student_class)

    if 'update_items' in request.POST:
        student = Student.objects.filter(pk=_as_int(request.POST.get('update_items'))).first()
        if student is None:
            messages.error(request, 'That student was not found.')
            return render(request, 'src/std-alts.html', {})
        scoped_class(request.user, student.student_class_id)
        new_class = scoped_class(request.user, request.POST.get('class'))
        name = request.POST.get('name', '').strip()
        if not name:
            messages.error(request, 'The student name cannot be empty.')
            return class_list(new_class)
        student.student_name = name
        student.student_class = new_class
        student.save()
        messages.success(request, 'Record Updated Successfully')
        return class_list(new_class)

    return render(request, 'src/std-alts.html', {})


# ------------------------------------------------------------------ deletion (soft, confirmed)

def _confirm_delete(request, template, queryset, sel, what, form_name):
    """Two-step delete: show what will go, then require a reason and the class name."""
    base = {'subjects': Subject.objects.all()}
    if refuse_if_locked(request, sel):
        return render(request, template, base)
    count = queryset.count()
    if not count:
        messages.error(request, f'There are no {what} for {selection_label(sel)}.')
        return render(request, template, base)
    confirm = dict(base, confirm=True, count=count, what=what, sel=sel, label=selection_label(sel),
                   form_name=form_name)
    if 'confirm_delete' not in request.POST:
        return render(request, template, confirm)
    typed = request.POST.get('confirm_class', '').strip().lower()
    reason = request.POST.get('reason', '').strip()
    if typed != sel['class'].name.strip().lower() or not reason:
        messages.error(request, 'Type the class name exactly and give a reason to confirm the delete.')
        return render(request, template, confirm)
    batch = uuid.uuid4().hex
    with transaction.atomic():
        ids = []
        for row in queryset:
            row.soft_delete(request.user, reason=reason, batch=batch)
            ids.append(row.id)
        audit.log('bulk_delete', model=queryset.model.__name__,
                  object_repr=f'{len(ids)} {what}: {selection_label(sel)}', section_id=sel['class'].section_id,
                  details={'batch': batch, 'reason': reason, 'ids': ids, 'class_id': sel['class'].id,
                           'session_id': sel['session'].id, 'term_id': sel['term'].id,
                           'subject_id': sel['subject'].id if sel.get('subject') else None})
    messages.info(request, f'Deleted {len(ids)} {what} for {selection_label(sel)}. '
                           'They can be restored from the Recycle Bin.')
    return render(request, template, base)


@require('delete_records')
def delete_result(request):
    if request.method == 'POST' and ('form1' in request.POST or 'form2' in request.POST):
        whole_class = 'form1' in request.POST
        sel, error = read_selection(request, need_subject=not whole_class, include_archived=True)
        if error:
            messages.error(request, error)
            return render(request, 'src/resultdelete.html', {'subjects': Subject.objects.all()})
        results = StudentResult.objects.filter(session=sel['session'], term=sel['term'], student_class=sel['class'])
        if not whole_class:
            results = results.filter(subject=sel['subject'])
        return _confirm_delete(request, 'src/resultdelete.html', results, sel, 'results',
                               'form1' if whole_class else 'form2')
    return render(request, 'src/resultdelete.html', {'subjects': Subject.objects.all()})


@require('delete_records')
def delete_bhv(request):
    if request.method == 'POST' and 'form1' in request.POST:
        sel, error = read_selection(request, need_subject=False, include_archived=True)
        if error:
            messages.error(request, error)
            return render(request, 'src/bhvdelete.html', {})
        rows = StudentBehaviouralAssessment.objects.filter(session=sel['session'], term=sel['term'],
                                                           student_class=sel['class'])
        return _confirm_delete(request, 'src/bhvdelete.html', rows, sel, 'behavioural assessments', 'form1')
    return render(request, 'src/bhvdelete.html', {})
